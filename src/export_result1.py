# -*- coding: utf-8 -*-
"""从已通过验证的 Q1 结果生成两个明确不同用途的 result1 工作簿（不重新求解）。

双版本策略
----------
* ``outputs/result1_submission.xlsx`` —— **官方模板保留版本**（submission-safe）。
  严格保留官方模板的全部已有时间标签、文字、格式、sheet 与结构，**只填写原本需要
  填写的数值单元格**（计划购电量 B2:B145、六个 4h 充电/放电量、E0、E144）。
  ``计划购电量!A2:A145`` **保持官方模板原样，绝不修改**——即使官方标签与模型
  Mapping A 相差一格。其时间标签来自官方模板，不代表重新定义模型时间映射。

* ``outputs/result1_corrected_mapping.xlsx`` —— **Mapping A 修正版本**。
  数值与提交版完全一致，但 ``计划购电量!A2:A145`` 显式修正为当天物理区间
  ``00:00-00:10 … 23:50-24:00``（见 ``docs/time_mapping_decision.md`` 第 3.3 节）。
  用于内部核对、论文、数据分析与 Mapping A 物理解释；**不默认作为正式提交文件**。

两份文件数值写入顺序均为 ``model_t = t`` → 电子表格第 ``t + 1`` 行（``output_row = t + 1``）。
两份都先在临时文件写出并重新读取验证，**两份全部通过后**才原子替换正式路径；
任何一份验证失败都不会发布损坏结果。官方模板源文件本身永不被修改。
所有输出数值单位均为 kWh。
"""
from __future__ import annotations

import json
import math
import os
import shutil
import sys
import tempfile
from copy import copy
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.model_q1 import DELTA_T, E_INIT, N_PERIODS  # noqa: E402

TEMPLATE = ROOT / "C题" / "附件" / "附件5" / "result1.xlsx"
SOLUTION = ROOT / "outputs" / "q1_solution.csv"
SUMMARY = ROOT / "outputs" / "q1_summary.json"
OUTPUT_SUBMISSION = ROOT / "outputs" / "result1_submission.xlsx"
OUTPUT_CORRECTED = ROOT / "outputs" / "result1_corrected_mapping.xlsx"
VALUE_TOL = 1e-8  # Excel/CSV floating-point serialization only; no rounding.
TOTAL_TOL = 1e-6
TEMP_PREFIX = ".result1-"
STYLE_ATTRIBUTES = ("font", "fill", "border", "alignment", "protection", "number_format")

# 官方模板“计划购电量”标签整体相对物理区间后移一格（见 time_mapping_decision §3.3）。
MAPPING_A_FIRST_INTERVAL = "00:00-00:10"
MAPPING_A_LAST_INTERVAL = "23:50-24:00"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(f"result1 validation failed: {message}")


def close(actual, expected, label: str, tol: float = TOTAL_TOL) -> None:
    require(isinstance(actual, (int, float)) and not isinstance(actual, bool)
            and isinstance(expected, (int, float)) and not isinstance(expected, bool)
            and math.isfinite(actual) and math.isfinite(expected)
            and abs(actual - expected) <= tol, label)


def validate_inputs(sol: pd.DataFrame, summary: dict) -> None:
    require(len(sol) == N_PERIODS, "expected 144 input rows")
    require(sol["model_t"].tolist() == list(range(1, N_PERIODS + 1)), "model_t order must be 1..144")
    require(sol["physical_interval"].astype(str).tolist()[0] == MAPPING_A_FIRST_INTERVAL
            and sol["physical_interval"].astype(str).tolist()[-1] == MAPPING_A_LAST_INTERVAL,
            "source CSV intervals must follow Mapping A (00:00-00:10 ... 23:50-24:00)")
    require(summary.get("validation_all_pass") is True and summary.get("status") == "OPTIMAL",
            "source summary must be validated and OPTIMAL")
    require(summary.get("n_periods") == N_PERIODS, "summary n_periods")
    close(summary.get("delta_t_hours"), DELTA_T, "summary delta_t_hours")
    for direction in ("grid", "charge", "discharge"):
        kw = sol[f"{direction}_kw"].to_numpy(float)
        kwh = sol[f"{direction}_kwh"].to_numpy(float)
        require(bool(np.isfinite(kw).all() and np.isfinite(kwh).all()), f"{direction} must be finite")
        require(bool((kw >= 0).all() and (kwh >= 0).all()), f"{direction} must be nonnegative")
        require(bool(np.allclose(kwh, kw * DELTA_T, rtol=0, atol=VALUE_TOL)), f"{direction}_kwh != kw * 1/6")
        close(float(kwh.sum()), summary.get(f"total_{direction}_energy_kwh"), f"total {direction}")
    close(float(sol["energy_start_kwh"].iloc[0]), E_INIT, "CSV E0")
    close(float(sol["energy_end_kwh"].iloc[-1]), E_INIT, "CSV E144")
    close(summary.get("initial_energy_kwh"), E_INIT, "summary E0")
    close(summary.get("final_energy_kwh"), E_INIT, "summary E144")
    close(float((sol["price_yuan_per_kwh"] * sol["grid_kwh"]).sum()),
          summary.get("objective_cost_yuan"), "objective reconciliation")
    blocks = summary["blocks_4h"]
    require(len(blocks) == 6, "expected six summary blocks")
    for i, block in enumerate(blocks):
        for direction in ("charge", "discharge"):
            close(float(sol[f"{direction}_kwh"].iloc[i * 24:(i + 1) * 24].sum()),
                  block[f"{direction}_kwh_sum"], f"block {i} {direction}")


# --------------------------------------------------------------------------
# 模板保真与数值校验
# --------------------------------------------------------------------------
def _assert_template_preserved(output_wb, template_wb, changed) -> None:
    """除 ``changed`` 判定的单元格（仅数值）外，输出必须与官方模板逐格一致。"""
    require(output_wb.sheetnames == template_wb.sheetnames, "official sheet names/order")
    for ws, source in zip(output_wb.worksheets, template_wb.worksheets):
        require(str(ws.merged_cells) == str(source.merged_cells), f"{ws.title!r} merged cells changed")
        require(ws.freeze_panes == source.freeze_panes, f"{ws.title!r} freeze panes changed")
        require(ws.max_row == source.max_row and ws.max_column == source.max_column,
                f"{ws.title!r} dimensions changed")
        for row in source.iter_rows():
            for cell in row:
                target = ws[cell.coordinate]
                for attribute in STYLE_ATTRIBUTES:
                    require(copy(getattr(target, attribute)) == copy(getattr(cell, attribute)),
                            f"{ws.title!r}!{cell.coordinate} style changed")
                if not changed(ws, cell):
                    require(target.value == cell.value,
                            f"{ws.title!r}!{cell.coordinate} value changed")


def _submission_changed(ws, cell) -> bool:
    """提交版允许变化的单元格：仅原本待填写的数值格，A 列永不改动。"""
    if ws.title == "计划购电量":
        return cell.row >= 2 and cell.column == 2          # 仅 B 列
    if ws.title == "充放电量":
        return (2 <= cell.row <= 7 and cell.column in (2, 3)) or cell.coordinate in ("E2", "E3")
    return False


def _corrected_changed(ws, cell) -> bool:
    """修正版：A、B 两列数据行 + 充放电量数值格。"""
    if ws.title == "计划购电量":
        return cell.row >= 2 and cell.column <= 2          # A、B 两列
    if ws.title == "充放电量":
        return (2 <= cell.row <= 7 and cell.column in (2, 3)) or cell.coordinate in ("E2", "E3")
    return False


def _validate_filled_values(wb, sol: pd.DataFrame, summary: dict) -> None:
    """逐行校验 B2:B145、六个 4h 块、E0/E144（两份版本共用）。"""
    grid, battery = wb.worksheets
    require(grid.max_row == 145, "expected exactly 144 purchase rows")
    for i, row in sol.iterrows():
        excel_row = i + 2
        value = grid.cell(excel_row, 2).value
        close(value, float(row["grid_kwh"]), f"row {excel_row} grid kWh", VALUE_TOL)
        require(value >= 0, f"row {excel_row} negative purchase")
    for i, block in enumerate(summary["blocks_4h"]):
        for col, direction in ((2, "charge"), (3, "discharge")):
            close(battery.cell(i + 2, col).value, block[f"{direction}_kwh_sum"],
                  f"block {i} {direction} output")
    close(battery["E2"].value, E_INIT, "output E0")
    close(battery["E3"].value, E_INIT, "output E144")


def validate_submission_workbook(path: Path, sol: pd.DataFrame, summary: dict,
                                 template: Path = TEMPLATE) -> None:
    """官方模板保留版：A 列/文字/格式/sheet/结构逐格等于模板，仅数值格被填写。"""
    validate_inputs(sol, summary)
    output_wb = openpyxl.load_workbook(path, data_only=False)
    template_wb = openpyxl.load_workbook(template, data_only=False)
    try:
        _assert_template_preserved(output_wb, template_wb, _submission_changed)
        grid = output_wb.worksheets[0]
        tgrid = template_wb.worksheets[0]
        for row in range(2, 146):
            require(grid.cell(row, 1).value == tgrid.cell(row, 1).value,
                    f"submission A{row} must equal official template label")
        _validate_filled_values(output_wb, sol, summary)
    finally:
        output_wb.close()
        template_wb.close()


def validate_corrected_workbook(path: Path, sol: pd.DataFrame, summary: dict,
                                template: Path = TEMPLATE) -> None:
    """Mapping A 修正版：A2/A145 及逐行标签为当天物理区间，且与模板仅标签+数值差异。"""
    validate_inputs(sol, summary)
    expected = sol["physical_interval"].astype(str).tolist()
    require(expected[0] == MAPPING_A_FIRST_INTERVAL
            and expected[-1] == MAPPING_A_LAST_INTERVAL, "source CSV is not Mapping A")
    output_wb = openpyxl.load_workbook(path, data_only=False)
    template_wb = openpyxl.load_workbook(template, data_only=False)
    try:
        require(output_wb.sheetnames == template_wb.sheetnames, "official sheet names/order")
        _assert_template_preserved(output_wb, template_wb, _corrected_changed)
        grid, battery = output_wb.worksheets
        tbattery = template_wb.worksheets[1]
        require(grid.max_row == 145, "expected exactly 144 purchase rows")
        require(grid["A2"].value == MAPPING_A_FIRST_INTERVAL, "corrected A2 must be 00:00-00:10")
        require(grid["A145"].value == MAPPING_A_LAST_INTERVAL, "corrected A145 must be 23:50-24:00")
        for i, row in sol.iterrows():
            excel_row = i + 2
            require(grid.cell(excel_row, 1).value == row["physical_interval"],
                    f"row {excel_row} interval")
            value = grid.cell(excel_row, 2).value
            require(value >= 0, f"row {excel_row} negative purchase")
            close(value, float(row["grid_kwh"]), f"row {excel_row} grid kWh", VALUE_TOL)
            close(value, float(row["grid_kw"]) * DELTA_T, f"row {excel_row} grid conversion", VALUE_TOL)
        for i in range(6):
            require(battery.cell(i + 2, 1).value == tbattery.cell(i + 2, 1).value,
                    f"block {i} template label")
        _validate_filled_values(output_wb, sol, summary)
    finally:
        output_wb.close()
        template_wb.close()


# 向后兼容别名：旧名 ``validate_workbook`` 指 Mapping A 修正版。
validate_workbook = validate_corrected_workbook


# --------------------------------------------------------------------------
# 写盘
# --------------------------------------------------------------------------
def _fill_values(wb, sol: pd.DataFrame, summary: dict) -> None:
    grid, battery = wb.worksheets
    for i, row in sol.iterrows():
        grid.cell(i + 2, 2, float(row["grid_kwh"]))
    for i, block in enumerate(summary["blocks_4h"]):
        battery.cell(i + 2, 2, block["charge_kwh_sum"])
        battery.cell(i + 2, 3, block["discharge_kwh_sum"])
    battery["E2"] = summary["initial_energy_kwh"]
    battery["E3"] = summary["final_energy_kwh"]


def _write_submission(temp_path: Path, template: Path, sol: pd.DataFrame, summary: dict) -> None:
    shutil.copyfile(template, temp_path)
    wb = openpyxl.load_workbook(temp_path)
    try:
        require(wb.sheetnames == ["计划购电量", "充放电量"], "unexpected template sheets")
        require(wb.worksheets[0]["A1"].value == "时间段"
                and wb.worksheets[0]["B1"].value == "购电量", "template purchase headers")
        _fill_values(wb, sol, summary)   # 不动 A 列
        wb.save(temp_path)
    finally:
        wb.close()


def _write_corrected(temp_path: Path, template: Path, sol: pd.DataFrame, summary: dict) -> None:
    shutil.copyfile(template, temp_path)
    wb = openpyxl.load_workbook(temp_path)
    try:
        require(wb.sheetnames == ["计划购电量", "充放电量"], "unexpected template sheets")
        grid = wb.worksheets[0]
        require(grid["A1"].value == "时间段" and grid["B1"].value == "购电量",
                "template purchase headers")
        for i, row in sol.iterrows():
            grid.cell(i + 2, 1, row["physical_interval"])   # Mapping A 修正标签
        _fill_values(wb, sol, summary)
        wb.save(temp_path)
    finally:
        wb.close()


def _make_temp(directory: Path) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    fd, name = tempfile.mkstemp(prefix=TEMP_PREFIX, suffix=".xlsx", dir=directory)
    os.close(fd)
    return Path(name)


def export_result1(
    solution: Path = SOLUTION, summary: Path = SUMMARY, template: Path = TEMPLATE,
    *, submission_output: Path | None = OUTPUT_SUBMISSION,
    corrected_output: Path = OUTPUT_CORRECTED, output: Path | None = None,
) -> tuple[Path, Path | None]:
    """生成双版本 result1；两份全部验证通过后才原子替换。

    返回 ``(corrected_path, submission_path)``。若传入旧参数 ``output``，则仅生成
    Mapping A 修正版到该路径（submission 为 ``None``），用于向后兼容与单文件检查。
    """
    solution, summary, template = map(Path, (solution, summary, template))
    if output is not None:
        corrected_output = Path(output)
        submission_output = None
    else:
        corrected_output = Path(corrected_output)
        submission_output = None if submission_output is None else Path(submission_output)

    source_resolved = {solution.resolve(), summary.resolve(), template.resolve(), TEMPLATE.resolve()}
    targets = [corrected_output] + ([submission_output] if submission_output is not None else [])
    for target in targets:
        require(target.resolve() not in source_resolved,
                "output must not overwrite a source or the official template")
    require(len({t.resolve() for t in targets}) == len(targets), "output paths must be distinct")

    sol = pd.read_csv(solution, encoding="utf-8-sig").reset_index(drop=True)
    summ = json.loads(summary.read_text(encoding="utf-8"))
    validate_inputs(sol, summ)

    created: dict[Path, Path] = {}
    temps: list[Path] = []
    try:
        temp = _make_temp(corrected_output.parent)
        temps.append(temp)
        _write_corrected(temp, template, sol, summ)
        validate_corrected_workbook(temp, sol, summ, template)   # 发布前验证
        created[corrected_output] = temp

        if submission_output is not None:
            temp = _make_temp(submission_output.parent)
            temps.append(temp)
            _write_submission(temp, template, sol, summ)
            validate_submission_workbook(temp, sol, summ, template)  # 发布前验证
            created[submission_output] = temp

        # 两份全部验证通过 → 原子替换
        for target, temp_path in created.items():
            os.replace(temp_path, target)
    finally:
        for temp_path in temps:
            Path(temp_path).unlink(missing_ok=True)
    return corrected_output, submission_output


if __name__ == "__main__":
    corrected, submission = export_result1()
    print(f"PASS (submission) : {submission}")
    print("  official template preserved; A2:A145 kept verbatim; values only in B2:B145 + blocks + E0/E144.")
    print(f"PASS (corrected)  : {corrected}")
    print("  Mapping A: A2:A145 = 00:00-00:10 ... 23:50-24:00; output_row = model_t + 1.")
