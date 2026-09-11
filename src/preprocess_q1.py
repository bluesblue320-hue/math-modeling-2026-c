# -*- coding: utf-8 -*-
"""
A2 Data Engineering Agent — Problem 1 输入数据标准化与验证。

职责边界（严格遵守 AGENTS.md 第 4 节 Scope）：
  * 只读原始附件 1，统一时间解析，按 Mapping A 建立 model_t，
    生成 physical_interval，执行数据完整性验证，输出标准化数据集。
  * 不建模、不求解、不写 result1.xlsx、不模拟储能、不生成 P_grid/P_ch/P_dis/E_t。

唯一事实源：
  * 时间映射：docs/time_mapping_decision.md  (STATUS: DECIDED, Mapping A)
  * 题目事实：docs/problem_spec.md
  * 数学模型：docs/model_spec.md

允许 hardcode 的内容（来自 Source of Truth，非数据）：
  * DELTA_T_H = 1/6 h（10 min 步长）
  * 时间映射规则（source_row = model_t + 1，右端点解释）
  * 列名/单位定义

禁止 hardcode：任何 price / load / pv 数值 —— 必须真实从 Excel 读取。
"""

from __future__ import annotations

import datetime as _dt
import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd

# --------------------------------------------------------------------------
# 常量（源自 Source of Truth，允许 hardcode）
# --------------------------------------------------------------------------
DELTA_T_H: float = 1.0 / 6.0          # 时间步长 = 10 min，model_spec.md §1
N_PERIODS: int = 144                  # 一天 144 个时段
SHEET_NAME: str = "Sheet1"            # problem_spec.md §2.1
HEADER_ROW: int = 1                   # 第 1 行表头
FIRST_SOURCE_ROW: int = 2             # 第 2 行起为数据
LAST_SOURCE_ROW: int = 145            # 第 145 行结束
TIME_STEP_MIN: int = 10               # 相邻标签间隔固定 10 min
MINUTE_FULL_DAY: int = 1440           # 0:00+1 = 1440 min

# 列索引（附件 1 A1:D145）
COL_TIME = 1
COL_PRICE = 2
COL_LOAD = 3
COL_PV = 4

# 输出字段（题面指定）
OUTPUT_COLUMNS = [
    "model_t",
    "physical_interval",
    "interval_start_min",
    "interval_end_min",
    "source_excel_row",
    "source_label",
    "source_minute",
    "price_yuan_per_kwh",
    "load_kw",
    "pv_kw",
]

# 默认路径（基于本文件位置推导仓库根）
REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC_FILE = REPO_ROOT / "C题" / "附件" / "附件1.xlsx"
DEFAULT_OUT_CSV = REPO_ROOT / "data" / "processed" / "q1_data.csv"
DEFAULT_OUT_JSON = REPO_ROOT / "outputs" / "a2_validation.json"

# 题目表 1 六个关键时段（physical_interval -> model_t），Mapping A
# time_mapping_decision.md §5 / model_spec.md §8
KEY_MAPPING_CHECKS: list[dict[str, Any]] = [
    {"model_t": 1,   "physical_interval": "00:00-00:10", "source_excel_row": 2,   "source_label": "00:10"},
    {"model_t": 61,  "physical_interval": "10:00-10:10", "source_excel_row": 62,  "source_label": "10:10"},
    {"model_t": 73,  "physical_interval": "12:00-12:10", "source_excel_row": 74,  "source_label": "12:10"},
    {"model_t": 85,  "physical_interval": "14:00-14:10", "source_excel_row": 86,  "source_label": "14:10"},
    {"model_t": 97,  "physical_interval": "16:00-16:10", "source_excel_row": 98,  "source_label": "16:10"},
    {"model_t": 109, "physical_interval": "18:00-18:10", "source_excel_row": 110, "source_label": "18:10"},
    {"model_t": 121, "physical_interval": "20:00-20:10", "source_excel_row": 122, "source_label": "20:10"},
    {"model_t": 144, "physical_interval": "23:50-24:00", "source_excel_row": 145, "source_label": "0:00+1"},
]


# --------------------------------------------------------------------------
# 时间解析（关键：混合类型，禁用 astype(str) + sort_values）
# --------------------------------------------------------------------------
def parse_source_time(value: Any) -> int:
    """将附件 1 A 列单元格统一解析为 source_minute（分钟，范围 10..1440）。

    支持：
      * datetime.time / datetime.datetime  （第 2–61 行的 Excel 时间对象）
      * 字符串 "H:MM" / "HH:MM"           （第 62–144 行）
      * 字符串 "0:00+1"                    → 1440 min（次日 0:00 = 当天 24:00）

    其中 "0:00+1" 处理：字符串以 '+1' 结尾即代表次日 0:00，
    其相对当天 00:00 的分钟数为 1440，而非 0。
    """
    if value is None:
        raise ValueError("时间单元格为空，无法解析")

    # Excel 时间对象
    if isinstance(value, _dt.datetime):
        return value.hour * 60 + value.minute
    if isinstance(value, _dt.time):
        return value.hour * 60 + value.minute

    # 字符串
    if isinstance(value, str):
        s = value.strip()
        if not s:
            raise ValueError("时间单元格为空字符串，无法解析")
        # 次日标记："0:00+1" → 1440
        if s.endswith("+1"):
            base = s[:-2].strip()
            minute = _parse_hhmm(base)
            return minute + MINUTE_FULL_DAY
        return _parse_hhmm(s)

    raise ValueError(f"无法识别的时间类型: {type(value)!r} 值={value!r}")


def _parse_hhmm(s: str) -> int:
    """解析 'H:MM' 或 'HH:MM' 为分钟数。"""
    parts = s.split(":")
    if len(parts) != 2:
        raise ValueError(f"时间格式非法: {s!r}")
    hh, mm = parts
    if not (hh.strip().isdigit() and mm.strip().isdigit()):
        raise ValueError(f"时间格式非法: {s!r}")
    h = int(hh)
    m = int(mm)
    if not (0 <= m < 60):
        raise ValueError(f"分钟数越界: {s!r}")
    if not (0 <= h <= 24):
        raise ValueError(f"小时数越界: {s!r}")
    return h * 60 + m


def format_label(minute: int) -> str:
    """把 source_minute 规范化为展示标签。

    1440 → '0:00+1'（与附件 1 第 145 行原始写法一致）；其余 → 'HH:MM'。
    """
    if minute == MINUTE_FULL_DAY:
        return "0:00+1"
    h, m = divmod(minute, 60)
    return f"{h:02d}:{m:02d}"


def format_interval(start_min: int, end_min: int) -> str:
    """把物理区间 [start, end) 分钟格式化为 'HH:MM-HH:MM' 展示串。

    1440 写作 '24:00'（当天 24:00），不使用次日标记。
    """
    return f"{_hhmm24(start_min)}-{_hhmm24(end_min)}"


def _hhmm24(minute: int) -> str:
    h, m = divmod(minute, 60)
    return f"{h:02d}:{m:02d}"


# --------------------------------------------------------------------------
# 读取原始附件
# --------------------------------------------------------------------------
def read_raw_attachment(src_file: Path) -> list[dict[str, Any]]:
    """逐单元格读取附件 1（openpyxl，真实读取，不做任何插值/填充）。

    返回按行顺序排列的原始记录列表，每条含：
      excel_row, raw_time, raw_price, raw_load, raw_pv
    """
    if not src_file.exists():
        raise FileNotFoundError(f"附件 1 不存在: {src_file}")

    wb = openpyxl.load_workbook(src_file, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise ValueError(f"未找到工作表 {SHEET_NAME!r}，实际为 {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    records: list[dict[str, Any]] = []
    for r in range(FIRST_SOURCE_ROW, LAST_SOURCE_ROW + 1):
        records.append(
            {
                "excel_row": r,
                "raw_time": ws.cell(row=r, column=COL_TIME).value,
                "raw_price": ws.cell(row=r, column=COL_PRICE).value,
                "raw_load": ws.cell(row=r, column=COL_LOAD).value,
                "raw_pv": ws.cell(row=r, column=COL_PV).value,
            }
        )
    wb.close()
    return records


def header_info(src_file: Path) -> dict[str, Any]:
    """读取表头与原始维度，供 data_report.md 记录。"""
    wb = openpyxl.load_workbook(src_file, data_only=True, read_only=True)
    ws = wb[SHEET_NAME]
    header = [ws.cell(row=HEADER_ROW, column=c).value for c in range(1, COL_PV + 1)]
    dims = ws.calculate_dimension()
    max_row = ws.max_row
    max_col = ws.max_column
    wb.close()
    return {"header": header, "dimensions": dims, "max_row": max_row, "max_col": max_col}


# --------------------------------------------------------------------------
# 构建标准化数据集（Mapping A）
# --------------------------------------------------------------------------
def build_dataset(records: list[dict[str, Any]]) -> pd.DataFrame:
    """按 Mapping A 建立标准化数据集。

    Mapping A（右端点解释）：
      model_t = t (1..144)
      physical_interval = [10(t-1), 10t] min
      source_excel_row  = t + 1
      source_minute     = 10t
      source_label      = source_minute 的标签写法（1440 → '0:00+1'）
    """
    rows: list[dict[str, Any]] = []
    for idx, rec in enumerate(records):
        model_t = idx + 1
        interval_start_min = (model_t - 1) * TIME_STEP_MIN
        interval_end_min = model_t * TIME_STEP_MIN
        rows.append(
            {
                "model_t": model_t,
                "physical_interval": format_interval(interval_start_min, interval_end_min),
                "interval_start_min": interval_start_min,
                "interval_end_min": interval_end_min,
                "source_excel_row": rec["excel_row"],
                "source_label": format_label(parse_source_time(rec["raw_time"])),
                "source_minute": parse_source_time(rec["raw_time"]),
                "price_yuan_per_kwh": rec["raw_price"],
                "load_kw": rec["raw_load"],
                "pv_kw": rec["raw_pv"],
            }
        )
    df = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    return df


# --------------------------------------------------------------------------
# 验证（A–J）
# --------------------------------------------------------------------------
def validate(df: pd.DataFrame) -> dict[str, Any]:
    """执行全部完整性检查，返回 machine-readable 结果与 errors/warnings。"""
    errors: list[str] = []
    warnings: list[str] = []

    # A. 行数
    rows = int(len(df))
    if rows != N_PERIODS:
        errors.append(f"[A] 行数错误：期望 {N_PERIODS}，实际 {rows}")

    # B. model_t 严格 1..144
    expected_t = list(range(1, N_PERIODS + 1))
    actual_t = df["model_t"].tolist()
    duplicate_model_t = int(df["model_t"].duplicated().sum())
    if actual_t != expected_t:
        if sorted(actual_t) == expected_t:
            errors.append("[B] model_t 数值齐全但顺序错误")
        else:
            errors.append(f"[B] model_t 非严格 1..{N_PERIODS}：缺失/多余")

    # C. source_minute 严格 10..1440，diff=10
    expected_min = [i * TIME_STEP_MIN for i in range(1, N_PERIODS + 1)]
    actual_min = df["source_minute"].tolist()
    if actual_min != expected_min:
        errors.append(f"[C] source_minute 非严格 {TIME_STEP_MIN} 递增序列")
    diffs = pd.Series(actual_min).diff().dropna().unique().tolist()
    time_step_minutes = int(max(diffs)) if len(diffs) > 0 else 0
    if diffs != [TIME_STEP_MIN]:
        errors.append(f"[C] 相邻 source_minute 差不为常数 {TIME_STEP_MIN}：{diffs}")

    # D. physical interval 覆盖 00:00-24:00，无 gap/overlap
    starts = df["interval_start_min"].tolist()
    ends = df["interval_end_min"].tolist()
    if starts != expected_min_schedule(0):
        errors.append("[D] interval_start_min 不连续/有 gap 或 overlap")
    if ends != expected_min_schedule(TIME_STEP_MIN):
        errors.append("[D] interval_end_min 不连续/有 gap 或 overlap")
    if starts[0] != 0 or ends[-1] != MINUTE_FULL_DAY:
        errors.append("[D] physical interval 未覆盖 00:00-24:00")
    if any(e > MINUTE_FULL_DAY for e in ends):
        errors.append("[D] 出现次日额外区间（end > 1440）")
    duplicate_intervals = int(df["physical_interval"].duplicated().sum())

    # E. source_excel_row 必须 2..145
    expected_rows = list(range(FIRST_SOURCE_ROW, LAST_SOURCE_ROW + 1))
    actual_rows = df["source_excel_row"].tolist()
    if actual_rows != expected_rows:
        errors.append(f"[E] source_excel_row 非 {FIRST_SOURCE_ROW}..{LAST_SOURCE_ROW}")

    # F. 数值缺失
    value_cols = ["price_yuan_per_kwh", "load_kw", "pv_kw"]
    missing_values = int(df[value_cols].isna().sum().sum())
    if missing_values != 0:
        errors.append(f"[F] 存在 {missing_values} 个缺失值（不自动填补，需人工确认）")

    # G. 数值类型可转换
    numeric_cols: dict[str, pd.Series] = {}
    for col in value_cols:
        try:
            numeric_cols[col] = pd.to_numeric(df[col], errors="raise").astype(float)
        except (ValueError, TypeError) as exc:
            errors.append(f"[G] 列 {col} 无法转换为数值：{exc}")
            numeric_cols[col] = pd.Series([float("nan")] * rows)

    # H. 合理性检查（不修改，只报告）
    stats: dict[str, Any] = {}
    negative_price_count = 0
    negative_load_count = 0
    negative_pv_count = 0
    for col in value_cols:
        s = numeric_cols[col]
        s_min = float(s.min()) if s.notna().any() else float("nan")
        s_max = float(s.max()) if s.notna().any() else float("nan")
        s_mean = float(s.mean()) if s.notna().any() else float("nan")
        stats[col] = {"min": s_min, "max": s_max, "mean": s_mean}
        if s_min < 0:
            warnings.append(f"[H] {col} 存在负值：min={s_min}")
    negative_price_count = int((numeric_cols["price_yuan_per_kwh"] < 0).sum())
    negative_load_count = int((numeric_cols["load_kw"] < 0).sum())
    negative_pv_count = int((numeric_cols["pv_kw"] < 0).sum())

    # I. 重复
    duplicate_source_label = int(df["source_label"].duplicated().sum())
    if duplicate_model_t:
        errors.append(f"[I] model_t 重复 {duplicate_model_t} 个")
    if duplicate_intervals:
        errors.append(f"[I] physical_interval 重复 {duplicate_intervals} 个")
    if duplicate_source_label:
        errors.append(f"[I] source_label 重复 {duplicate_source_label} 个")

    # J. 时间关键点验证
    key_mapping: list[dict[str, Any]] = []
    for chk in KEY_MAPPING_CHECKS:
        t = chk["model_t"]
        sub = df[df["model_t"] == t]
        ok = False
        actual: dict[str, Any] = {}
        if len(sub) == 1:
            r = sub.iloc[0]
            actual = {
                "physical_interval": r["physical_interval"],
                "source_excel_row": int(r["source_excel_row"]),
                "source_label": r["source_label"],
                "source_minute": int(r["source_minute"]),
            }
            ok = (
                actual["physical_interval"] == chk["physical_interval"]
                and actual["source_excel_row"] == chk["source_excel_row"]
                and actual["source_label"] == chk["source_label"]
            )
        else:
            actual = {"error": f"model_t={t} 行数={len(sub)}"}
        if not ok:
            errors.append(f"[J] 关键点 model_t={t} 映射不符：期望 {chk}，实际 {actual}")
        key_mapping.append({"expected": chk, "actual": actual, "ok": bool(ok)})

    # 汇总
    status = "PASS" if not errors else "FAIL"
    result: dict[str, Any] = {
        "status": status,
        "rows": rows,
        "missing_values": missing_values,
        "duplicate_model_t": duplicate_model_t,
        "duplicate_intervals": duplicate_intervals,
        "duplicate_source_label": duplicate_source_label,
        "time_step_minutes": time_step_minutes,
        "first_model_t": int(df["model_t"].iloc[0]) if rows else None,
        "last_model_t": int(df["model_t"].iloc[-1]) if rows else None,
        "first_interval": str(df["physical_interval"].iloc[0]) if rows else None,
        "last_interval": str(df["physical_interval"].iloc[-1]) if rows else None,
        "first_source_row": int(df["source_excel_row"].iloc[0]) if rows else None,
        "last_source_row": int(df["source_excel_row"].iloc[-1]) if rows else None,
        "first_source_label": str(df["source_label"].iloc[0]) if rows else None,
        "last_source_label": str(df["source_label"].iloc[-1]) if rows else None,
        "full_day_covered": bool(starts[0] == 0 and ends[-1] == MINUTE_FULL_DAY) if rows else False,
        "mapping": "RIGHT_ENDPOINT",
        "mapping_decision": "Mapping A (source label = interval end)",
        "delta_t_hours": DELTA_T_H,
        "negative_price_count": negative_price_count,
        "negative_load_count": negative_load_count,
        "negative_pv_count": negative_pv_count,
        "statistics": stats,
        "key_mapping_checks": key_mapping,
        "errors": errors,
        "warnings": warnings,
    }
    return result


def expected_min_schedule(offset_min: int) -> list[int]:
    """生成期望的区间端点序列。"""
    return [i * TIME_STEP_MIN + offset_min for i in range(N_PERIODS)]


# --------------------------------------------------------------------------
# 主流程
# --------------------------------------------------------------------------
def describe_source_path(src_file: Path) -> str:
    """把输入文件路径序列化为**可移植**的写法。

    产物（validation JSON）不应写入本机绝对路径（例如 `C:\\Users\\...\\数模\\...`）。
    若文件位于仓库根目录之内，记为 POSIX 风格的仓库相对路径（如 `C题/附件/附件1.xlsx`）；
    否则退化为文件名。这样同一份产物在不同机器/克隆目录下生成结果一致。
    """
    try:
        return src_file.resolve().relative_to(REPO_ROOT.resolve()).as_posix()
    except ValueError:
        return src_file.name


def run_preprocess(
    src_file: Path = DEFAULT_SRC_FILE,
    out_csv: Path = DEFAULT_OUT_CSV,
    out_json: Path = DEFAULT_OUT_JSON,
    print_summary: bool = True,
) -> dict[str, Any]:
    """完整流程：读取 → 解析 → 建表 → 验证 → 输出。"""
    hdr = header_info(src_file)
    records = read_raw_attachment(src_file)
    df = build_dataset(records)
    result = validate(df)

    result["source_file"] = describe_source_path(src_file)
    result["sheet"] = SHEET_NAME
    result["header"] = hdr["header"]
    result["raw_dimensions"] = hdr["dimensions"]
    result["raw_max_row"] = hdr["max_row"]
    result["raw_max_col"] = hdr["max_col"]

    # 输出 CSV（无论 PASS / FAIL 都写，便于人工检查；STATUS 由 JSON 明确）
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_csv, index=False, encoding="utf-8-sig")

    # 输出 machine-readable validation JSON
    out_json.parent.mkdir(parents=True, exist_ok=True)
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    if print_summary:
        _print_summary(result, out_csv, out_json)

    return result


def _print_summary(result: dict[str, Any], out_csv: Path, out_json: Path) -> None:
    print("=" * 68)
    print(f"A2 DATA STATUS: {result['status']}")
    print("=" * 68)
    print(f"源文件            : {result.get('source_file')}")
    print(f"sheet / 原始维度  : {result.get('sheet')} / {result.get('raw_dimensions')}")
    print(f"原始行数          : {result.get('raw_max_row', 0) - 1} 数据行")
    print(f"标准化行数        : {result['rows']}")
    print(f"缺失值            : {result['missing_values']}")
    print(f"重复 model_t      : {result['duplicate_model_t']}")
    print(f"重复 interval     : {result['duplicate_intervals']}")
    print(f"重复 source_label : {result['duplicate_source_label']}")
    print(f"时间步长          : {result['time_step_minutes']} min")
    print(f"第一时段          : {result['first_interval']} "
          f"(model_t={result['first_model_t']}, row={result['first_source_row']}, "
          f"label={result['first_source_label']})")
    print(f"最后时段          : {result['last_interval']} "
          f"(model_t={result['last_model_t']}, row={result['last_source_row']}, "
          f"label={result['last_source_label']})")
    print(f"全天覆盖 00:00-24:00: {'YES' if result['full_day_covered'] else 'NO'}")
    print("-" * 68)
    print("关键映射点:")
    for km in result["key_mapping_checks"]:
        e, a = km["expected"], km["actual"]
        flag = "OK" if km["ok"] else "XX"
        print(f"  [{flag}] t={e['model_t']:>3} {e['physical_interval']:<11} "
              f"row={e['source_excel_row']:<4} label={e['source_label']:<7} | actual={a}")
    print("-" * 68)
    print("统计量:")
    for col, st in result["statistics"].items():
        print(f"  {col:<20} min={st['min']:.4f}  max={st['max']:.4f}  mean={st['mean']:.4f}")
    print(f"负值计数: price={result['negative_price_count']} "
          f"load={result['negative_load_count']} pv={result['negative_pv_count']}")
    if result["warnings"]:
        print("-" * 68)
        print("WARNINGS:")
        for w in result["warnings"]:
            print(f"  ! {w}")
    if result["errors"]:
        print("-" * 68)
        print("ERRORS:")
        for e in result["errors"]:
            print(f"  X {e}")
    print("-" * 68)
    print(f"CSV : {out_csv}")
    print(f"JSON: {out_json}")
    print("=" * 68)


def main() -> int:
    result = run_preprocess()
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    sys.exit(main())
