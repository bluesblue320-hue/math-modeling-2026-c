"""Finalization regression tests: fail-fast ``_val``, dual result1 export, paper assets, A4.

覆盖本轮 11 项测试要求；所有测试均复用已冻结的 validated outputs，不重新清洗数据。
"""
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pulp
import pytest

from src import build_q1_paper_assets as paper
from src import solve_q1
from src.export_result1 import (
    SOLUTION,
    SUMMARY,
    TEMPLATE,
    export_result1,
    validate_corrected_workbook,
    validate_submission_workbook,
)
from src.model_q1 import build_model

REPO_ROOT = Path(__file__).resolve().parent.parent
MAIN_OBJECTIVE = 35126.948589289634
BASELINE_OBJECTIVE = 48052.046590846665
SAVING_OBJECTIVE = 12925.098001557031
SAVING_RATE = 0.26898121762870153


def _solved():
    df = solve_q1.load_data()
    inputs = [df[c].tolist() for c in ("price_yuan_per_kwh", "load_kw", "pv_kw")]
    bundle = build_model(*inputs)
    solve_q1.solve_stage(bundle, stage_label="finalize")
    return df, bundle


def _export_pair(tmp_path):
    original = TEMPLATE.read_bytes()
    corrected, submission = export_result1(
        submission_output=tmp_path / "result1_submission.xlsx",
        corrected_output=tmp_path / "result1_corrected_mapping.xlsx",
    )
    assert TEMPLATE.read_bytes() == original, "official template source file must not change"
    return corrected, submission


# 1. _val(None) 必须 fail-fast
def test_val_none_raises_fail_fast():
    variable = pulp.LpVariable("probe_without_value")
    assert variable.value() is None
    with pytest.raises(RuntimeError, match="has no value"):
        solve_q1._val(variable)


# 2. 正常 _val() 不改变原结果
def test_val_normal_does_not_change_results():
    df, bundle = _solved()
    assert solve_q1._val(bundle.p_grid[1]) == float(bundle.p_grid[1].value())
    sol = solve_q1.extract_solution(df, bundle)
    actual = solve_q1.build_solution_dataframe(df, sol)
    frozen = pd.read_csv(SOLUTION)
    assert actual["physical_interval"].tolist() == frozen["physical_interval"].tolist()
    np.testing.assert_allclose(actual.select_dtypes("number"), frozen.select_dtypes("number"),
                               rtol=0, atol=1e-8)
    assert float(pulp.value(bundle.cost_expr)) == pytest.approx(MAIN_OBJECTIVE, rel=0, abs=1e-6)


# 3. submission 版 A 列与官方模板完全一致，且模板源文件未变
def test_submission_keeps_official_template_labels(tmp_path):
    _, submission = _export_pair(tmp_path)
    template = openpyxl.load_workbook(TEMPLATE)
    output = openpyxl.load_workbook(submission)
    try:
        assert output.sheetnames == template.sheetnames
        tgrid, sgrid = template["计划购电量"], output["计划购电量"]
        for row in range(1, 146):
            assert sgrid.cell(row, 1).value == tgrid.cell(row, 1).value
        assert sgrid["A2"].value == "0:10-0:20"
        assert sgrid["A145"].value == "0:00+1-0:10+1"
        # 充放电量表的表头与文字标签（非数值格）保持模板原样
        batt, tbatt = output["充放电量"], template["充放电量"]
        assert [batt.cell(1, c).value for c in range(1, 6)] == \
               [tbatt.cell(1, c).value for c in range(1, 6)]
        for row in range(2, 8):
            assert batt.cell(row, 1).value == tbatt.cell(row, 1).value
        assert batt["D2"].value == tbatt["D2"].value
        assert batt["D3"].value == tbatt["D3"].value
    finally:
        template.close()
        output.close()
    validate_submission_workbook(submission, pd.read_csv(SOLUTION), json.loads(SUMMARY.read_text()))


# 4. corrected 版使用 Mapping A 标签
def test_corrected_uses_mapping_a_labels(tmp_path):
    corrected, _ = _export_pair(tmp_path)
    wb = openpyxl.load_workbook(corrected)
    try:
        grid = wb["计划购电量"]
        assert grid["A2"].value == "00:00-00:10"
        assert grid["A62"].value == "10:00-10:10"
        assert grid["A145"].value == "23:50-24:00"
    finally:
        wb.close()
    validate_corrected_workbook(corrected, pd.read_csv(SOLUTION), json.loads(SUMMARY.read_text()))


# 5. 两份 Excel 数值结果一致，且等于 validated output
def test_two_versions_numerically_identical(tmp_path):
    corrected, submission = _export_pair(tmp_path)
    sol = pd.read_csv(SOLUTION)
    cw = openpyxl.load_workbook(corrected)
    sw = openpyxl.load_workbook(submission)
    try:
        cgrid, cbatt = cw.worksheets
        sgrid, sbatt = sw.worksheets
        grid_matrix = [sgrid.cell(r, 2).value for r in range(2, 146)]
        assert [cgrid.cell(r, 2).value for r in range(2, 146)] == grid_matrix
        np.testing.assert_allclose(grid_matrix, sol["grid_kwh"].to_numpy(float), rtol=0, atol=1e-8)
        for row in range(2, 8):
            assert cbatt.cell(row, 2).value == sbatt.cell(row, 2).value
            assert cbatt.cell(row, 3).value == sbatt.cell(row, 3).value
        for cell in ("E2", "E3"):
            assert cbatt[cell].value == sbatt[cell].value == 6000.0
    finally:
        cw.close()
        sw.close()


# 6. paper table1 数值来自 validated output
def test_paper_table1_matches_validated_output(tmp_path, monkeypatch):
    sol, summ = paper.load_validated()
    rows = paper.build_table1(sol, summ)
    assert [r["时间段"] for r in rows] == paper.TABLE1_INTERVALS
    by_interval = {str(r.physical_interval): float(r.grid_kwh) for r in sol.itertuples()}
    for row in rows:
        assert row["购电量(kWh)"] == pytest.approx(by_interval[row["时间段"]], rel=0, abs=1e-9)
    # 关键值抽查（来自 validated q1_summary.json 的 key_intervals）
    key = {k["physical_interval"]: k["grid_kwh"] for k in summ["key_intervals"]}
    for row in rows:
        assert row["购电量(kWh)"] == pytest.approx(key[row["时间段"]], rel=0, abs=1e-6)
    # 写出后的 CSV 与内存结果一致
    monkeypatch.setattr(paper, "PAPER_DIR", tmp_path / "paper")
    monkeypatch.setattr(paper, "FIGURES_DIR", tmp_path / "paper" / "figures")
    paper.write_tables(sol, summ)
    text = (tmp_path / "paper" / "q1_table1.csv").read_text(encoding="utf-8-sig")
    for row in rows:
        assert f"{row['时间段']},{row['购电量(kWh)']}" in text


# 7. paper table2 4h 汇总正确
def test_paper_table2_blocks_correct():
    sol, summ = paper.load_validated()
    rows = paper.build_table2(summ)
    assert [r["时间段"] for r in rows] == [b["block"] for b in summ["blocks_4h"]]
    for row, block in zip(rows, summ["blocks_4h"]):
        assert row["充电量(kWh)"] == pytest.approx(block["charge_kwh_sum"], rel=0, abs=1e-9)
        assert row["放电量(kWh)"] == pytest.approx(block["discharge_kwh_sum"], rel=0, abs=1e-9)
    # 与原始逐时段切片独立核对
    charge = sol["charge_kwh"].to_numpy(float)
    discharge = sol["discharge_kwh"].to_numpy(float)
    for i, row in enumerate(rows):
        assert row["充电量(kWh)"] == pytest.approx(charge[i * 24:(i + 1) * 24].sum(), rel=0, abs=1e-6)
        assert row["放电量(kWh)"] == pytest.approx(discharge[i * 24:(i + 1) * 24].sum(), rel=0, abs=1e-6)
    energy = paper.build_table2_energy(sol)
    assert [e["储电量(kWh)"] for e in energy] == [6000.0, 6000.0]


# 8. paper metrics 与 q1_summary 一致
def test_paper_metrics_match_summary():
    _, summ = paper.load_validated()
    metrics = paper.build_metrics(summ)
    assert metrics["objective_cost_yuan"] == pytest.approx(MAIN_OBJECTIVE, rel=0, abs=1e-6)
    assert metrics["baseline_cost_yuan"] == pytest.approx(BASELINE_OBJECTIVE, rel=0, abs=1e-6)
    assert metrics["saving_yuan"] == pytest.approx(SAVING_OBJECTIVE, rel=0, abs=1e-6)
    assert metrics["saving_rate"] == pytest.approx(SAVING_RATE, rel=0, abs=1e-12)
    assert metrics["objective_cost_yuan"] == summ["objective_cost_yuan"]
    assert metrics["baseline_cost_yuan"] == summ["baseline_no_storage"]["baseline_cost_yuan"]
    assert metrics["saving_yuan"] == summ["baseline_no_storage"]["saving_yuan"]
    assert metrics["saving_rate"] == summ["baseline_no_storage"]["saving_rate"]
    assert metrics["pv_curtailment_kwh"] == summ["total_curtailment_kwh"] == 0.0
    assert metrics["efficiency_main"] == {"eta_ch": 0.9, "eta_dis": 0.9, "round_trip": 0.81}


# 9. figure generation 正常完成
def test_figure_generation_completes(tmp_path, monkeypatch):
    monkeypatch.setattr(paper, "PAPER_DIR", tmp_path / "paper")
    monkeypatch.setattr(paper, "FIGURES_DIR", tmp_path / "paper" / "figures")
    result = paper.run()
    assert (tmp_path / "paper" / "q1_table1.csv").exists()
    assert (tmp_path / "paper" / "q1_table2.csv").exists()
    assert (tmp_path / "paper" / "q1_paper_metrics.json").exists()
    names = {p.name for p in result["figures"]}
    for required in ("q1_power_dispatch.png", "q1_storage_charge_discharge.png",
                     "q1_storage_energy.png", "q1_cost_comparison.png"):
        assert required in names
    for path in result["figures"]:
        assert path.exists() and path.stat().st_size > 0
    metrics = json.loads((tmp_path / "paper" / "q1_paper_metrics.json").read_text(encoding="utf-8"))
    assert metrics["objective_cost_yuan"] == pytest.approx(MAIN_OBJECTIVE, rel=0, abs=1e-6)


# 论文素材已提交产物存在（交付就绪）
def test_committed_paper_artifacts_exist():
    for name in ("q1_table1.csv", "q1_table2.csv", "q1_paper_metrics.json"):
        assert (paper.PAPER_DIR / name).exists(), f"missing committed paper artifact {name}"
    for name in ("q1_power_dispatch.png", "q1_storage_charge_discharge.png",
                 "q1_storage_energy.png", "q1_cost_comparison.png"):
        assert (paper.FIGURES_DIR / name).exists(), f"missing committed figure {name}"


# 11. A4 independent audit 继续 PASS
def test_a4_independent_audit_passes():
    proc = subprocess.run(
        [sys.executable, str(REPO_ROOT / "src" / "audit_q1_independent.py")],
        cwd=REPO_ROOT, capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert "A4 STATUS: PASS" in proc.stdout
    assert "35126.948589289634" in proc.stdout
