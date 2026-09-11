"""Regression, parameter plumbing, export corruption and preservation tests."""
import json
import math
from copy import copy

import numpy as np
import openpyxl
import pandas as pd
import pulp
import pytest

from src import solve_q1
from src.model_q1 import ETA_CH, ETA_DIS, DELTA_T, build_model
from src.sensitivity_q1_efficiency import run as run_sensitivity
from src.export_result1 import (
    export_result1, validate_workbook, TEMPLATE, SOLUTION, SUMMARY,
)


def solved(eta=None):
    df = solve_q1.load_data()
    inputs = [df[c].tolist() for c in ("price_yuan_per_kwh", "load_kw", "pv_kw")]
    bundle = build_model(*inputs, **({} if eta is None else {"eta_ch": eta, "eta_dis": eta}))
    solve_q1.solve_stage(bundle, stage_label="test")
    return df, bundle, solve_q1.extract_solution(df, bundle)


def test_default_efficiency_and_frozen_dispatch():
    df, bundle, sol = solved()
    assert bundle.eta_ch == ETA_CH == 0.90
    assert bundle.eta_dis == ETA_DIS == 0.90
    recursion = bundle.prob.constraints["C3_recursion_1"]
    assert recursion[bundle.p_ch[1]] == pytest.approx(-0.90 * DELTA_T)
    assert recursion[bundle.p_dis[1]] == pytest.approx(DELTA_T / 0.90)
    assert pulp.value(bundle.cost_expr) == pytest.approx(35126.948589289634, rel=0, abs=1e-6)
    actual = solve_q1.build_solution_dataframe(df, sol)
    frozen = pd.read_csv(SOLUTION)
    assert actual["physical_interval"].tolist() == frozen["physical_interval"].tolist()
    np.testing.assert_allclose(actual.select_dtypes("number"), frozen.select_dtypes("number"), rtol=0, atol=1e-8)


def test_custom_efficiency_validation_uses_actual_parameters():
    eta = math.sqrt(0.90)
    df, bundle, sol = solved(eta)
    cost = pulp.value(bundle.cost_expr)
    good = solve_q1.run_validation(df, sol, cost, eta_ch=bundle.eta_ch, eta_dis=bundle.eta_dis)
    assert good["all_checks_pass"]
    assert good["simultaneous_charge_discharge_count"] == 0
    assert not solve_q1.run_validation(df, sol, cost)["checks"]["energy_recursion"]
    assert ETA_CH == ETA_DIS == 0.90


@pytest.mark.parametrize("eta", [0, -0.1, 1.1, float("nan"), float("inf")])
@pytest.mark.parametrize("direction", ["eta_ch", "eta_dis"])
def test_invalid_efficiency_rejected(eta, direction):
    with pytest.raises(ValueError, match=direction):
        build_model([1.] * 144, [4000.] * 144, [0.] * 144, **{direction: eta})


def test_sensitivity_outputs_and_main_files_preserved(tmp_path):
    before = [p.read_bytes() for p in (SOLUTION, SUMMARY)]
    payload = run_sensitivity(tmp_path)
    saved = json.loads((tmp_path / "q1_efficiency_sensitivity.json").read_text())
    csv = pd.read_csv(tmp_path / "q1_efficiency_sensitivity.csv")
    assert payload == saved
    assert csv["scenario"].tolist() == ["A", "B"]
    a, b = saved["scenarios"]
    assert a["round_trip_efficiency"] == pytest.approx(0.81)
    assert b["round_trip_efficiency"] == pytest.approx(0.90)
    assert b["objective_cost_yuan"] < a["objective_cost_yuan"]
    assert saved["comparison"]["objective_B_minus_A_yuan"] == b["objective_cost_yuan"] - a["objective_cost_yuan"]
    assert all(s["validation_all_pass"] for s in (a, b))
    assert before == [p.read_bytes() for p in (SOLUTION, SUMMARY)]


def test_export_and_template_preservation(tmp_path):
    original = TEMPLATE.read_bytes()
    path = export_result1(output=tmp_path / "result1.xlsx")
    sol = pd.read_csv(SOLUTION)
    summ = json.loads(SUMMARY.read_text())
    validate_workbook(path, sol, summ)
    assert TEMPLATE.read_bytes() == original
    wb, template = openpyxl.load_workbook(path), openpyxl.load_workbook(TEMPLATE)
    try:
        assert wb["计划购电量"].max_row == 145
        assert wb["计划购电量"]["A2"].value == "00:00-00:10"
        assert wb["计划购电量"]["A145"].value == "23:50-24:00"
        for ws, source in zip(wb, template):
            assert str(ws.merged_cells) == str(source.merged_cells)
            assert ws.freeze_panes == source.freeze_panes
            for row in source:
                for cell in row:
                    target = ws[cell.coordinate]
                    # A serialized default StyleArray and absent style are equivalent.
                    for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
                        assert copy(getattr(target, attribute)) == copy(getattr(cell, attribute))
                    changed = (ws.title == "计划购电量" and cell.row >= 2 and cell.column <= 2) or (
                        ws.title == "充放电量" and ((2 <= cell.row <= 7 and cell.column in (2, 3))
                                                 or cell.coordinate in ("E2", "E3")))
                    if not changed:
                        assert ws[cell.coordinate].value == cell.value
    finally:
        wb.close()
        template.close()


@pytest.mark.parametrize("sheet,cell,value", [
    ("计划购电量", "A2", "00:10-00:20"),
    ("计划购电量", "A145", "24:00-24:10"),
    ("计划购电量", "B2", -1),
    ("计划购电量", "B62", 100),
    ("计划购电量", "A146", "extra"),
    ("充放电量", "B2", 0),
    ("充放电量", "C3", 0),
    ("充放电量", "E2", 5999),
    ("充放电量", "E3", 5999),
])
def test_export_validation_detects_corruption(tmp_path, sheet, cell, value):
    path = export_result1(output=tmp_path / "result1.xlsx")
    wb = openpyxl.load_workbook(path)
    wb[sheet][cell] = value
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="validation failed"):
        validate_workbook(path, pd.read_csv(SOLUTION), json.loads(SUMMARY.read_text()))


@pytest.mark.parametrize("corruption", ["unit", "nan", "order", "block", "failed"])
def test_bad_sources_do_not_publish(tmp_path, corruption):
    sol, summ = pd.read_csv(SOLUTION), json.loads(SUMMARY.read_text())
    if corruption == "unit":
        sol.loc[0, "grid_kwh"] = sol.loc[0, "grid_kw"]
    elif corruption == "nan":
        sol.loc[0, "grid_kwh"] = float("nan")
    elif corruption == "order":
        sol = sol.iloc[::-1]
    elif corruption == "block":
        summ["blocks_4h"][0]["charge_kwh_sum"] += 10
    else:
        summ["validation_all_pass"] = False
    source, summary, output = tmp_path / "s.csv", tmp_path / "s.json", tmp_path / "result1.xlsx"
    sol.to_csv(source, index=False)
    summary.write_text(json.dumps(summ))
    with pytest.raises(ValueError):
        export_result1(source, summary, output=output)
    assert not output.exists()


def test_failed_workbook_validation_keeps_previous_valid_output(tmp_path, monkeypatch):
    import src.export_result1 as exporter
    path = export_result1(output=tmp_path / "result1.xlsx")
    before = path.read_bytes()
    def fail(*args):
        raise ValueError("injected saved workbook validation failure")
    monkeypatch.setattr(exporter, "validate_workbook", fail)
    with pytest.raises(ValueError, match="injected"):
        export_result1(output=path)
    assert path.read_bytes() == before
    assert list(tmp_path.glob(".result1-*.xlsx")) == []


def test_representative_optimum_separate_output(tmp_path, monkeypatch):
    before = [p.read_bytes() for p in (SOLUTION, SUMMARY)]
    monkeypatch.setattr(solve_q1, "REPO_ROOT", tmp_path)
    result = solve_q1.run(print_summary=False, representative_optimum=True)
    secondary = result["secondary_optimization"]
    assert result["secondary_optimization_used"]
    assert result["validation_all_pass"]
    assert secondary["cost_diff"] <= solve_q1.EPS_COST + 1e-8
    assert secondary["throughput_after_kwh"] <= secondary["throughput_before_kwh"] + 1e-6
    assert (tmp_path / "outputs/representative_optimum/q1_solution.csv").exists()
    assert before == [p.read_bytes() for p in (SOLUTION, SUMMARY)]


def test_failed_main_validation_does_not_write(monkeypatch):
    original_validation = solve_q1.run_validation
    def fail(*args, **kwargs):
        result = original_validation(*args, **kwargs)
        result["all_checks_pass"] = False
        result["checks"]["power_balance"] = False
        return result
    writes = []
    monkeypatch.setattr(solve_q1, "run_validation", fail)
    monkeypatch.setattr(solve_q1, "write_outputs", lambda *args, **kwargs: writes.append(True))
    with pytest.raises(RuntimeError, match="内部验证未全部通过"):
        solve_q1.run(print_summary=False)
    assert writes == []
